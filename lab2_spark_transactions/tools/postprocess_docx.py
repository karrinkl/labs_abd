#!/usr/bin/env python3
"""Доводит report.docx после pandoc до вида учебного отчёта.

- разрыв страницы после титульного листа и содержание (поле TOC, заполненное
  заголовками; номера страниц берутся из --pages, рассчитанного tools/toc_pages.py);
- рамки таблиц, шрифт 10 pt в таблицах, повтор строки заголовка;
- подписи «Таблица N — ...» над таблицами, без абзацного отступа;
- нумерация страниц внизу по центру (на титульном листе номер не ставится);
- «Приложение А» с новой страницы.

Запуск: python3 tools/postprocess_docx.py report.docx [--pages pages.json]
"""
import argparse
import json
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm, Pt


def add_field(paragraph, instr, placeholder=""):
    """Вставляет сложное поле Word (PAGE, TOC и т. п.)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run._r.append(begin)
    run = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    run._r.append(it)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    if placeholder:
        paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    tbl_pr.append(borders)
    # ширина таблицы — 100 % ширины текста
    for old in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(old)
    w = OxmlElement("w:tblW")
    w.set(qn("w:type"), "pct")
    w.set(qn("w:w"), "5000")
    tbl_pr.append(w)


def format_table(table):
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        if r_idx == 0:
            tr_pr = row._tr.get_or_add_trPr()
            hdr = OxmlElement("w:tblHeader")
            tr_pr.append(hdr)
        for cell in row.cells:
            for p in cell.paragraphs:
                pf = p.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.line_spacing = 1.0
                pf.space_before = Pt(1)
                pf.space_after = Pt(1)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(10)
                    if r_idx == 0:
                        run.font.bold = True


def fld(run, kind):
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    run._r.append(el)


def add_toc(doc, before, headings, pages):
    """Поле TOC, результат которого заранее заполнен заголовками и номерами страниц.

    Word и LibreOffice могут пересчитать его штатно (F9 / «Обновить все»).
    """
    text_width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for i, (style, text) in enumerate(headings):
        p = before.insert_paragraph_before()
        p.style = doc.styles["Normal"]
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0 if style == "Heading 1" else 0.75)
        pf.line_spacing = 1.15
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if i == 0:
            fld(p.add_run(), "begin")
            r = p.add_run()
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = 'TOC \\o "1-2" \\h \\z \\u'
            r._r.append(it)
            fld(p.add_run(), "separate")
        p.add_run(f"{text}\t{pages.get(text, '')}")
        if i == len(headings) - 1:
            fld(p.add_run(), "end")


M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def merge_math_runs(body):
    """Склеивает соседние прямые (m:sty=p) буквы формулы: pandoc пишет «CR» как
    «C» + «R», и LibreOffice показывает их с пробелом. Word результат не меняет."""
    def upright_letters(r):
        if r is None or r.tag != f"{{{M_NS}}}r":
            return None
        sty = r.find(f"{{{M_NS}}}rPr/{{{M_NS}}}sty")
        t = r.find(f"{{{M_NS}}}t")
        if sty is None or sty.get(f"{{{M_NS}}}val") != "p" or t is None or not (t.text or "").isalpha():
            return None
        return t
    for r in list(body.iter(f"{{{M_NS}}}r")):
        t = upright_letters(r)
        if t is None or r.getparent() is None:
            continue
        nxt = r.getnext()
        while (nt := upright_letters(nxt)) is not None:
            t.text += nt.text
            after = nxt.getnext()
            nxt.getparent().remove(nxt)
            nxt = after


def main(path, pages):
    doc = Document(path)
    body = doc.element.body

    merge_math_runs(body)

    # Таблицы
    for t in doc.tables:
        format_table(t)

    paragraphs = doc.paragraphs
    first_h1 = None
    for p in paragraphs:
        style = p.style.name if p.style is not None else ""
        text = p.text.strip()
        if style == "Heading 1" and first_h1 is None and text:
            first_h1 = p
        if style == "Heading 1" and text.startswith("Приложение"):
            p.paragraph_format.page_break_before = True
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
        if re.match(r"^Таблица \d+ — ", text):
            pf = p.paragraph_format
            pf.first_line_indent = Cm(0)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.keep_with_next = True
            pf.space_before = Pt(6)
            pf.space_after = Pt(3)
        if text.startswith("[МЕСТО ДЛЯ СКРИНШОТА") or text.startswith("[ЗАПОЛНИТЬ"):
            for run in p.runs:
                run.font.highlight_color = 7  # жёлтый маркер: место для заполнения
        if style == "Source Code":
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                run.font.size = Pt(8.5)
        if style == "Image Caption":
            p.paragraph_format.keep_together = True

    # Список источников — без выравнивания по ширине (длинные URL)
    in_refs = False
    for p in paragraphs:
        name = p.style.name if p.style is not None else ""
        if name == "Heading 1":
            in_refs = "Список использованных источников" in p.text
        elif in_refs:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Отступ перед абзацем, следующим сразу за таблицей
    for tbl in body.iterchildren(qn("w:tbl")):
        nxt = tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            ppr = nxt.get_or_add_pPr()
            sp = ppr.find(qn("w:spacing"))
            if sp is None:
                sp = OxmlElement("w:spacing")
                ppr.append(sp)
            sp.set(qn("w:before"), "160")

    # Вертикальные отступы титульного листа
    seen = set()
    for p in paragraphs:
        name = p.style.name if p.style is not None else ""
        if name in ("Title Page Main", "Title Page Right", "Title Page Bottom") and name not in seen:
            seen.add(name)
            p.paragraph_format.space_before = Pt(
                {"Title Page Main": 90, "Title Page Right": 60, "Title Page Bottom": 120}[name]
            )

    # Содержание после титульного листа
    headings = [(p.style.name, p.text.strip()) for p in paragraphs
                if p.style is not None and p.style.name in ("Heading 1", "Heading 2")]
    if first_h1 is not None:
        toc_title = first_h1.insert_paragraph_before("СОДЕРЖАНИЕ")
        toc_title.style = doc.styles["Normal"]
        toc_title.paragraph_format.page_break_before = True
        toc_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        toc_title.paragraph_format.first_line_indent = Cm(0)
        toc_title.paragraph_format.space_after = Pt(12)
        toc_title.runs[0].font.bold = True
        add_toc(doc, first_h1, headings, pages)
        first_h1.paragraph_format.page_break_before = True

    # Нумерация страниц
    for section in doc.sections:
        section.different_first_page_header_footer = True
        fp = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.first_line_indent = Cm(0)
        add_field(fp, "PAGE", "1")

    # Попросить Word обновить поля при открытии
    settings = doc.settings.element
    upd = settings.find(qn("w:updateFields"))
    if upd is None:
        upd = OxmlElement("w:updateFields")
        settings.append(upd)
    upd.set(qn("w:val"), "true")

    doc.core_properties.title = "Лабораторная работа № 2. Распределённые вычисления в Apache Spark (вариант № 2)"
    doc.core_properties.author = "[ФИО студента]"
    doc.save(path)
    print(f"post-processed {path}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("docx", nargs="?", default="report.docx")
    ap.add_argument("--pages", help="JSON {заголовок: номер страницы} из tools/toc_pages.py")
    a = ap.parse_args()
    pg = {}
    if a.pages:
        with open(a.pages, encoding="utf-8") as fh:
            pg = json.load(fh)
    main(a.docx, pg)
